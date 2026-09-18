# -*- coding: utf-8 -*-
"""Bilan chiffre des quatre campagnes, calcule a partir des fichiers du disque.

Aucun chiffre n'est saisi a la main : tout est recompte depuis le classeur
Excel (base RPPS + statuts) et les deux journaux. Ecrit bilan_chiffre.md dans
chaque dossier de campagne, et 00_COMMUN/bilan_global.md.

    python3 outils/02_bilan.py
"""
from __future__ import annotations

import glob
import os
import sys
from collections import Counter
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config as C
import journal as J


def lire_base(campagne: str):
    motif = os.path.join(C.RACINE, campagne, "base_rpps_*.xlsx")
    fichiers = sorted(glob.glob(motif))
    if not fichiers:
        return None, None
    from openpyxl import load_workbook
    wb = load_workbook(fichiers[0], read_only=True, data_only=True)
    ws = wb["Base RPPS"] if "Base RPPS" in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    entetes = list(next(it))
    lignes = [dict(zip(entetes, r)) for r in it if any(v is not None for v in r)]
    wb.close()
    return os.path.basename(fichiers[0]), lignes


def bilan_campagne(campagne: str) -> tuple[str, dict]:
    nom_base, lignes = lire_base(campagne)
    recherches = J.lire_journal(campagne, "recherches")
    envois = J.lire_journal(campagne, "envois")
    conf = C.CAMPAGNES[campagne]

    total = len(lignes) if lignes is not None else 0
    statuts = Counter((l.get("Statut_Connect") or "NON_VERIFIE") for l in (lignes or []))
    verifies = total - statuts.get("NON_VERIFIE", 0)
    deps = Counter((l.get("Departement") or "?") for l in (lignes or []))

    chiffres = {
        "total": total,
        "verifies": verifies,
        "exact": statuts.get("EXACT", 0),
        "a_confirmer": statuts.get("A_CONFIRMER", 0),
        "ambigu": statuts.get("AMBIGU", 0),
        "absent": statuts.get("ABSENT", 0),
        "deja_contacte": statuts.get("DEJA_CONTACTE", 0),
        "exclu": statuts.get("EXCLU", 0),
        "non_verifie": statuts.get("NON_VERIFIE", 0),
        "recherches": len(recherches),
        "envois": len(envois),
    }

    ok_r, msg_r = J.verifier_integrite(campagne, "recherches")
    ok_e, msg_e = J.verifier_integrite(campagne, "envois")

    L = []
    L.append("# Bilan chiffre - %s" % conf["libelle"])
    L.append("")
    L.append("Specialite ciblee : %s" % conf["libelle_long"])
    L.append("Perimetre : Ile-de-France (75, 77, 78, 91, 92, 93, 94, 95)")
    L.append("Genere automatiquement le %s par `outils/02_bilan.py`." % date.today().isoformat())
    L.append("Base de reference : %s" % (nom_base or "AUCUNE (extraction RPPS non encore realisee)"))
    L.append("")
    L.append("| Indicateur | Nombre |")
    L.append("|---|---:|")
    L.append("| Total recense (professionnels uniques apres deduplication RPPS) | %d |" % chiffres["total"])
    L.append("| Profils verifies sur Doctolib Connect | %d |" % chiffres["verifies"])
    L.append("| Correspondances exactes (EXACT) | %d |" % chiffres["exact"])
    L.append("| Profils a confirmer (A_CONFIRMER) | %d |" % chiffres["a_confirmer"])
    L.append("| Cas ambigus / homonymes (AMBIGU) | %d |" % chiffres["ambigu"])
    L.append("| Profils indisponibles ou hors Connect (ABSENT) | %d |" % chiffres["absent"])
    L.append("| Deja contactes (DEJA_CONTACTE) | %d |" % chiffres["deja_contacte"])
    L.append("| Exclusions nominatives (EXCLU) | %d |" % chiffres["exclu"])
    L.append("| Restant a verifier (NON_VERIFIE) | %d |" % chiffres["non_verifie"])
    L.append("| Recherches Connect consignees | %d |" % chiffres["recherches"])
    L.append("| **Messages reellement envoyes** | **%d** |" % chiffres["envois"])
    L.append("")
    if deps:
        L.append("## Repartition par departement")
        L.append("")
        L.append("| Dep. | Libelle | Nombre |")
        L.append("|---|---|---:|")
        for d in sorted(deps):
            L.append("| %s | %s | %d |" % (d, C.DEPARTEMENTS_IDF.get(d, "?"), deps[d]))
        L.append("")
    L.append("## Integrite des journaux")
    L.append("")
    L.append("- Journal des recherches : %s %s" % ("INTACT" if ok_r else "ALERTE", msg_r))
    L.append("- Journal des envois : %s %s" % ("INTACT" if ok_e else "ALERTE", msg_e))
    L.append("")
    if chiffres["total"] == 0:
        L.append("> **Base vide.** L'extraction RPPS n'a pas encore pu etre realisee.")
        L.append("> Voir `00_COMMUN/BLOCAGES.md`. Les chiffres ci-dessus sont a zero")
        L.append("> parce qu'aucune donnee n'a ete collectee, et non parce qu'aucun")
        L.append("> professionnel ne correspond a la cible.")
    return "\n".join(L) + "\n", chiffres


def main() -> int:
    globaux = Counter()
    resume = []
    for campagne in C.CAMPAGNES:
        texte, chiffres = bilan_campagne(campagne)
        dossier = os.path.join(C.RACINE, campagne)
        os.makedirs(dossier, exist_ok=True)
        with open(os.path.join(dossier, "bilan_chiffre.md"), "w", encoding="utf-8") as f:
            f.write(texte)
        globaux.update(chiffres)
        resume.append((campagne, chiffres))
        print("[%s] total=%d verifies=%d exact=%d ambigu=%d envois=%d"
              % (campagne, chiffres["total"], chiffres["verifies"],
                 chiffres["exact"], chiffres["ambigu"], chiffres["envois"]))

    L = ["# Bilan global des quatre campagnes", "",
         "Genere le %s par `outils/02_bilan.py`." % date.today().isoformat(), "",
         "| Campagne | Recenses | Verifies | Exacts | A confirmer | Ambigus | Absents | Envoyes |",
         "|---|---:|---:|---:|---:|---:|---:|---:|"]
    for campagne, c in resume:
        L.append("| %s | %d | %d | %d | %d | %d | %d | %d |"
                 % (C.CAMPAGNES[campagne]["libelle"], c["total"], c["verifies"],
                    c["exact"], c["a_confirmer"], c["ambigu"], c["absent"], c["envois"]))
    L.append("| **TOTAL** | **%d** | **%d** | **%d** | **%d** | **%d** | **%d** | **%d** |"
             % (globaux["total"], globaux["verifies"], globaux["exact"],
                globaux["a_confirmer"], globaux["ambigu"], globaux["absent"],
                globaux["envois"]))
    L.append("")
    exclusions = J.charger_exclusions()
    L.append("Exclusions nominatives enregistrees : %d." % len(set(exclusions.values() or [])))
    L.append("")
    with open(os.path.join(C.DOSSIER_COMMUN, "bilan_global.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print("-> 00_COMMUN/bilan_global.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
