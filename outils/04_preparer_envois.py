# -*- coding: utf-8 -*-
"""Etape 4 : preparation des envois — tout sauf le clic et la lecture de la confirmation.

Ce script NE SE CONNECTE A RIEN et n envoie rien. Il prepare un carnet d envoi :
pour chaque professionnel autorise, un fichier texte brut pret a etre colle en
UNE SEULE operation, dans un ordre stable, avec la liste de controle a derouler.

Ce qui est automatise ici :
  - selection des seuls destinataires autorises (statut EXACT, hors exclusions,
    jamais contactes, toutes campagnes confondues) ;
  - verrouillage du texte sur la version exactement validee par le Dr Benjoar ;
  - production d un fichier .txt par destinataire, en texte brut, pret au collage ;
  - feuille de saisie des confirmations, a reverser ensuite dans le journal.

Ce qui reste humain, et doit le rester :
  - ouvrir la conversation et lire l historique du destinataire ;
  - coller et envoyer ;
  - lire ce que l interface affiche reellement.

    python3 outils/04_preparer_envois.py --campagne 03_RHUMATOLOGUES --version RHUM-v1
    python3 outils/04_preparer_envois.py --campagne 03_RHUMATOLOGUES --version RHUM-v1 --limite 25
"""
from __future__ import annotations

import argparse
import csv
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config as C
import journal as J


def lire_base(campagne: str):
    fichiers = sorted(glob.glob(os.path.join(C.RACINE, campagne, "base_rpps_*.xlsx")))
    if not fichiers:
        return None, []
    from openpyxl import load_workbook
    wb = load_workbook(fichiers[0], read_only=True, data_only=True)
    ws = wb["Base RPPS"] if "Base RPPS" in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    entetes = list(next(it))
    lignes = [dict(zip(entetes, r)) for r in it if any(v is not None for v in r)]
    wb.close()
    return os.path.basename(fichiers[0]), lignes


def nom_fichier_sur(rpps: str, nom: str, prenom: str, rang: int) -> str:
    base = C.normaliser("%s_%s" % (nom, prenom)).replace(" ", "_")
    base = re.sub(r"[^a-z0-9_]", "", base)[:40]
    return "%03d_%s_%s.txt" % (rang, base or "sans_nom", rpps or "sans_rpps")


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--campagne", required=True, choices=list(C.CAMPAGNES))
    p.add_argument("--version", required=True,
                   help="version de message validee (voir 03_valider_modele.py --lister)")
    p.add_argument("--limite", type=int, default=0,
                   help="ne preparer que les N premiers (0 = tous)")
    a = p.parse_args()

    modele = J.modele_valide(a.campagne, a.version)
    if modele is None:
        print("REFUS : la version %r n est pas enregistree comme validee pour %s."
              % (a.version, a.campagne), file=sys.stderr)
        print("Aucun envoi ne peut etre prepare avant validation ecrite du texte "
              "par le Dr Benjoar.", file=sys.stderr)
        return 1

    chemin_texte = os.path.join(C.RACINE, modele["fichier"])
    if not os.path.exists(chemin_texte):
        print("REFUS : le fichier du modele valide est introuvable : %s"
              % modele["fichier"], file=sys.stderr)
        return 1
    with open(chemin_texte, encoding="utf-8") as f:
        texte = f.read()
    if J.empreinte_texte(texte) != modele["empreinte_texte"]:
        print("REFUS : le fichier du modele a ete modifie depuis sa validation.\n"
              "Son empreinte ne correspond plus. Faites revalider le nouveau texte "
              "sous une NOUVELLE version.", file=sys.stderr)
        return 1

    nom_base, lignes = lire_base(a.campagne)
    if not lignes:
        print("Aucune base RPPS pour %s. Rien a preparer." % a.campagne, file=sys.stderr)
        print("Voir 00_COMMUN/BLOCAGES.md.", file=sys.stderr)
        return 1

    dossier = os.path.join(C.RACINE, a.campagne, "envois_a_effectuer_%s" % a.version)
    os.makedirs(dossier, exist_ok=True)

    retenus, ecartes = [], []
    for l in lignes:
        rpps = str(l.get("RPPS") or "").strip()
        nom = str(l.get("Nom") or "").strip()
        prenom = str(l.get("Prenom") or "").strip()
        statut = str(l.get("Statut_Connect") or "NON_VERIFIE").strip()
        profil = str(l.get("Connect_profil_url") or "").strip()
        autorise, motif = J.peut_envoyer(a.campagne, rpps, nom, prenom, statut, profil)
        (retenus if autorise else ecartes).append((l, motif))

    if a.limite:
        retenus = retenus[:a.limite]

    # Un fichier texte brut par destinataire : jamais de saisie manuelle.
    feuille = os.path.join(dossier, "_feuille_de_confirmation.csv")
    with open(feuille, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rang", "fichier_texte", "rpps", "nom", "prenom", "ville",
                    "profil_connect", "historique_verifie_OUI",
                    "confirmation_interface", "remarque"])
        for rang, (l, _) in enumerate(retenus, start=1):
            rpps = str(l.get("RPPS") or "").strip()
            nom = str(l.get("Nom") or "").strip()
            prenom = str(l.get("Prenom") or "").strip()
            fichier = nom_fichier_sur(rpps, nom, prenom, rang)
            with open(os.path.join(dossier, fichier), "w", encoding="utf-8",
                      newline="\n") as ft:
                ft.write(texte)
            w.writerow([rang, fichier, rpps, nom, prenom,
                        str(l.get("Ville") or ""), str(l.get("Connect_profil_url") or ""),
                        "", "", ""])

    with open(os.path.join(dossier, "_MODE_OPERATOIRE.md"), "w", encoding="utf-8") as f:
        f.write("""# Mode operatoire — %s, version %s

%d destinataire(s) prepare(s). **Rien n'a ete envoye.**

Le texte est identique dans tous les fichiers : c'est la version validee %s
(empreinte %s), au caractere pres.

## Pour chaque ligne de `_feuille_de_confirmation.csv`, dans l'ordre

1. Ouvrir le profil du destinataire sur Doctolib Connect.
2. **Ouvrir la conversation et lire son historique.** Si un echange anterieur
   existe, ne pas envoyer : noter la remarque et passer au suivant.
3. Cocher `historique_verifie_OUI` avec `oui`.
4. Ouvrir le fichier `.txt` correspondant, **tout selectionner, copier**.
5. **Coller en une seule operation** dans la zone de message. Ne jamais retaper :
   une saisie caractere par caractere enverrait chaque retour a la ligne comme
   un message distinct.
6. Verifier de visu que le message colle est complet avant d'envoyer.
7. Envoyer, puis **lire ce que l'interface affiche reellement** et le recopier
   dans `confirmation_interface` (par exemple : « message affiche dans le fil,
   horodate 10:12 »). Laisser vide si rien n'est confirme.

## En cas d'erreur de saisie, de navigation ou de destinataire

Ne pas recommencer tout de suite. Verifier dans cet ordre : la conversation
(un message est-il parti ?), le brouillon (un residu a effacer ?), l'historique
du destinataire. Noter l'incident en remarque avant toute reprise.

## Une fois la feuille remplie

    python3 outils/05_consigner_envois.py --campagne %s --version %s \\
        --feuille %s

Seules les lignes portant une confirmation d'interface non vide seront
consignees comme envoyees. Les autres ne le seront pas.
""" % (C.CAMPAGNES[a.campagne]["libelle"], a.version, len(retenus), a.version,
       modele["empreinte_texte"][:16], a.campagne, a.version,
       os.path.relpath(feuille, C.RACINE)))

    print("Campagne  : %s" % C.CAMPAGNES[a.campagne]["libelle"])
    print("Base      : %s (%d lignes)" % (nom_base, len(lignes)))
    print("Version   : %s (empreinte %s)" % (a.version, modele["empreinte_texte"][:16]))
    print("Prepares  : %d destinataire(s) -> %s"
          % (len(retenus), os.path.relpath(dossier, C.RACINE)))
    print("Ecartes   : %d" % len(ecartes))
    from collections import Counter
    for motif, n in Counter(m.split(" :")[0].split(".")[0] for _, m in ecartes).most_common():
        print("    %4d  %s" % (n, motif))
    print("\nAUCUN MESSAGE N A ETE ENVOYE. Deroulez _MODE_OPERATOIRE.md.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
